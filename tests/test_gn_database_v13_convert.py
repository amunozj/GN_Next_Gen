"""Unit tests for the v1.3 converter's two tricky reads.

Both are exercised on a workbook built in the test, so the 28 MB artifact is not a
test fixture:

* ``Differences with HS98`` is four changelogs side by side in one sheet, not one
  table (columns A/I/Q/Y), and has to be unstacked;
* the ``Observers`` sheet marks v1.3 changes only as cell fill, which no cell value
  records (known issue D20).
"""

import datetime as dt
import importlib.util
import sys
from pathlib import Path

import pandas as pd
import pytest

openpyxl = pytest.importorskip("openpyxl")

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "convert_gn_database_v13.py"


@pytest.fixture(scope="module")
def conv():
    spec = importlib.util.spec_from_file_location("convert_gn_database_v13", SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = mod
    spec.loader.exec_module(mod)
    return mod


def _workbook(path: Path) -> None:
    from openpyxl.styles import PatternFill

    wb = openpyxl.Workbook()
    green = PatternFill(start_color="FF92D050", end_color="FF92D050", fill_type="solid")

    first = wb.active
    first.title = "1610-1946"
    first.append(["YEAR", "MONTH", "DAY", "STATION", "OBSER", "GROUPS"])
    first.append([1610, 1, 1, 0, 0, -1])
    first.append([1611, 3, 4, 2, 1, 3])
    first.append([None, None, None, None, None, None])  # trailing blank row

    second = wb.create_sheet("1947-1988")
    second.append(["YEAR", "MONTH", "DAY", "STATION", "OBSER", "GROUPS"])
    second.append([1947, 1, 1, 332, 1, 7])

    third = wb.create_sheet("1989-2010")
    third.append(["YEAR", "MONTH", "DAY", "STATION", "OBSER", "GROUPS"])
    third.append([2010, 12, 31, 445, 1, 2])

    diff = wb.create_sheet("Differences with HS98")
    diff.append(
        ["v1.12"]
        + [None] * 7
        + ["v1.2"]
        + [None] * 7
        + ["v1.21"]
        + [None] * 7
        + ["v1.3"]
        + [None] * 2
    )
    diff.append(
        ["NAME", "YEAR", "NOTES"]
        + [None] * 5
        + ["NAME", "YEAR", "NOTES"]
        + [None] * 5
        + ["NAME", "YEAR", "NOTES"]
        + [None] * 5
        + ["NAME", "YEAR", "NOTES"]
    )
    row = [None] * 27
    row[0:3] = ["OLD, A.", 1612, "old note"]
    row[8:11] = ["MID, B.", 1616, "mid note"]
    row[24:27] = ["NEW, C.", "1610-1613", "new note with doi"]
    diff.append(row)
    row2 = [None] * 27
    row2[24:27] = ["NEW, D.", 1611, "second new note"]
    diff.append(row2)

    obsv = wb.create_sheet("Observers")
    obsv.append(
        [
            "STATION",
            "INITIAL",
            "FINAL",
            "TOT.OBS",
            "OBSERVER",
            None,
            None,
            None,
            "All new changes are indicated in green",
        ]
    )
    obsv.append([1, 1610, 1613, 214, "HARRIOT, T., OXFORD"])
    obsv.append([2, "NA", "NA", 0, "MARIUS, S., NUREMBERG"])
    obsv.append([3, 1947, 1988, 900, "CHANGED, X."])
    obsv["A4"].fill = green  # only station 3 is flagged

    wb.save(path)


@pytest.fixture(scope="module")
def book(tmp_path_factory):
    p = tmp_path_factory.mktemp("v13") / "book.xlsx"
    _workbook(p)
    return p


def test_observations_concatenate_and_derive(conv, book):
    obs = conv.read_observations(book)
    assert len(obs) == 4, "the trailing blank row must be dropped"
    assert obs.groupby("SHEET").size().to_dict() == {"1610-1946": 2, "1947-1988": 1, "1989-2010": 1}
    first = obs.iloc[0]
    assert int(first["ORDINAL"]) == dt.date(1610, 1, 1).toordinal() == 587_676
    assert first["FRACYEAR"] == pytest.approx(1610.0)
    # Last day of a non-leap year: (365 - 1) / 365 past Jan 1.
    last = obs[obs["YEAR"] == 2010].iloc[0]
    assert last["FRACYEAR"] == pytest.approx(2010 + 364 / 365)
    assert list(obs.columns) == [
        "YEAR",
        "MONTH",
        "DAY",
        "ORDINAL",
        "STATION",
        "OBSER",
        "GROUPS",
        "FRACYEAR",
        "SHEET",
    ]


def test_observers_na_sentinel_and_fill(conv, book):
    df = conv.read_observers(book)
    assert len(df) == 3
    assert df.loc[df["STATION"] == 2, "INITIAL"].isna().all(), "'NA' must become null (D18)"
    assert df.loc[df["STATION"] == 2, "TOT_OBS"].iloc[0] == 0
    assert df.loc[df["STATION"] == 1, "INITIAL"].iloc[0] == 1610
    # Only the filled row is flagged -- the information is in formatting only (D20).
    assert df.set_index("STATION")["changed_in_v13"].to_dict() == {1: False, 2: False, 3: True}


def test_changelog_unstacks_four_blocks(conv, book):
    df = conv.read_changelog(book)
    assert df["version"].value_counts().to_dict() == {"v1.3": 2, "v1.12": 1, "v1.2": 1}
    assert "v1.21" not in set(df["version"]), "an empty block contributes no rows"
    v13 = df[df["version"] == "v1.3"]
    assert list(v13["name"]) == ["NEW, C.", "NEW, D."]
    # Years are kept as text because the sheet mixes ints and ranges like '1610-1613'.
    assert list(v13["year"]) == ["1610-1613", "1611"]


def test_parquet_write_is_deterministic(conv, tmp_path):
    df = pd.DataFrame({"a": [1, 2, 3], "b": ["x", "y", "z"]})
    one, two = tmp_path / "one.parquet", tmp_path / "two.parquet"
    conv.write_parquet(df, one)
    conv.write_parquet(df, two)
    assert conv.sha256(one) == conv.sha256(two)
    # No pandas metadata, so the bytes do not carry a library version.
    import pyarrow.parquet as pq

    assert pq.ParquetFile(one).schema_arrow.metadata in (None, {})
